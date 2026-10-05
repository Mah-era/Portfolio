"""Build the editable site copy inventory from the live local source data."""
import json
import subprocess
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
NODE = '/Users/apple/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node'
data = json.loads(subprocess.check_output([
    NODE, '--experimental-strip-types', '-e',
    "import('./lib/portfolio-data.ts').then(m => console.log(JSON.stringify(m)))",
], cwd=ROOT, text=True))
rows = []

def section(name):
    rows.append({'section': name})

def add(position, *texts):
    rows.append({'position': position, 'text': '\n'.join(str(t) for t in texts)})

section('Entrance and shared navigation')
add('Header / monogram and identity', 'mt', 'MAHERA TASFEE', 'A MIND FOR WHAT’S NEXT')
add('Header / navigation buttons', 'Profile', 'Projects 15', 'The floor plan')
add('Header / contact link', 'Let’s connect')
add('Entrance / eyebrow and edition', 'WELCOME TO MY WORLD', 'VOL. 01 / 2026')
add('Entrance / main heading', 'Mahera', 'Tasfee.')
add('Entrance / subheading', 'A business mind.', 'A builder’s instinct.')
add('Entrance / introduction', 'I connect the dots between people, data, and operations — and turn that thinking into things that work.')
add('Entrance / subject tags', 'Supply chain / Operations / Analytics')
add('Entrance / primary button', 'Step inside my world')
add('Entrance / secondary button', 'Or, go straight to my work')
add('Entrance / location and availability footnote', 'BASED IN DHAKA, BANGLADESH', 'OPEN TO OPPORTUNITIES')
add('Entrance / scene label', 'THE RESIDENCE')
add('Entrance / rotating seal', 'SCROLL TO EXPLORE · A WORLD OF IDEAS ·')
add('Entrance / architecture caption', 'AN EXPLORATION IN SEVEN CHAPTERS', 'Every room,', 'a different perspective.')
add('Entrance / scroll invitation', 'SCROLL TO ENTER', 'Take a little look around')

section('Shared room interface and display states')
add('Scene / day and night button states', 'GOLDEN HOUR', 'AFTER HOURS')
add('Scene / coordinates and caption states currently hidden by styling', '23.81° N / 90.41° E', *[f'CHAPTER {i:02d} / 07' for i in range(1, 8)], 'EVENING STUDY', 'LIGHT & PERSPECTIVE')
add('Scene / loading and fallback messages', 'Preparing the residence…', 'Opening the front door…', 'Explore the portfolio through the room guide.')
add('Scene / entrance and interior instructions', 'A scroll becomes a journey.', 'Drag to look around · scroll to walk · select an exhibit')
add('Room / collection label', 'THE COLLECTION')
add('Room / discover button states', 'Discover this room', 'Explore all 15 projects')
add('Footer / room journey labels', 'SCROLL TO CONTINUE', 'THE LOUNGE')
add('Footer / visited room counter states', *[f'{i} of 7 rooms explored' for i in range(1, 8)])
add('Footer / visibility controls', 'Clear view', 'Show labels')
add('Footer / camera control', 'Reset view')
add('Footer / motion control states', 'Motion', 'Still')

section('Room introductions and navigation labels')
for n, room_id in enumerate(data['roomOrder'], 1):
    room = data['roomData'][room_id]
    prefix = f'Room {n:02d} / {room["label"]}'
    add(prefix + ' / name and category in navigation and room sign', room['label'], room['category'])
    add(prefix + ' / chapter eyebrow', f'{n:02d} / 07', room['category'])
    add(prefix + ' / heading', room['title'])
    add(prefix + ' / subcopy and collection introduction', room['subtitle'])
    add(prefix + ' / story count', f'{len(data["roomItems"][room_id])} stories to discover')

section('Floor plan dialog')
add('Floor plan / eyebrow', 'YOUR VISIT, YOUR PACE')
add('Floor plan / title', 'A place for every', 'part of the story.')
add('Floor plan / description', 'Move directly to a room, or follow the full walkthrough.')
add('Floor plan / map annotation', 'N ↑', 'THE ENTRANCE', 'SEVEN ROOMS. ONE CURIOUS MIND.')
for n, room_id in enumerate(data['roomOrder'], 1):
    room = data['roomData'][room_id]
    add(f'Floor plan / room {n:02d} button and map tile', f'{n:02d}', room['label'], room['category'])
add('Floor plan / return and dismissal controls', 'Back to the front door', 'Close')

section('Exhibit and project interface')
add('3D exhibit / default and hover labels', 'A CLOSER LOOK ↗', 'OPEN STORY ↗')
add('3D exhibit / numbering in room collections', ', '.join(f'{i:02d}' for i in range(1, 16)))
add('3D architecture / entrance plaque', 'mt.', 'THE RESIDENCE')
add('3D architecture / room exit labels', *[data['roomData'][r]['label'].upper() for r in data['roomOrder'][1:]], 'UNTIL NEXT TIME')
add('3D architecture / portal button', 'CONTINUE →')
add('Project collection / introduction', 'The complete project collection. Open any project for a closer look.')
add('Project collection / thumbnail placeholder', 'SCREENSHOTS & FILM COMING SOON')
add('Collection / item button', 'Open story')
add('Project detail / media placeholder', 'Space for a screen recording, screenshots,', 'and the story behind this project.')
add('Project detail / media state labels', 'Project preview', 'Media coming later')
add('Exhibit detail / default external link label', 'Open link')
add('Exhibit detail / return button', 'See all stories in this room')
add('Collection and detail / dismissal control', 'Close')

for n, room_id in enumerate(data['roomOrder'], 1):
    room = data['roomData'][room_id]
    section(f'Room {n:02d} {room["label"]} exhibit copy')
    for j, item in enumerate(data['roomItems'][room_id], 1):
        fields = ['Eyebrow', 'Title', 'Description']
        texts = [item['eyebrow'], item['title'], item['body']]
        if item.get('language'):
            fields.append('Technology')
            texts.append(item['language'])
        if item.get('linkLabel'):
            fields.append('Link label')
            texts.append(item['linkLabel'])
        add(f'{room["label"]} / exhibit {j:02d}\n3D exhibit and collection and detail dialog\n' + ' / '.join(fields), *texts)

section('Accessibility and alternate content')
add('Header / home button accessible name', 'Mahera Tasfee, return to entrance')
add('Header / navigation accessible name', 'Portfolio shortcuts')
add('Day and night button / accessible names', 'Switch to daylight', 'Switch to evening')
add('Entrance / seal and arrow accessible names', 'Begin the seven-room journey', 'Enter the residence')
add('Footer / room navigation accessible name', 'Jump to room')
for n, room_id in enumerate(data['roomOrder'], 1):
    label = data['roomData'][room_id]['label']
    add(f'Room {n:02d} / navigation and map accessible names', f'{n}. {label}', f'Enter {label}')
add('Footer / visibility accessible names', 'Show portfolio labels', 'Hide labels for an unobstructed view')
add('Footer / camera accessible name', 'Recenter camera facing forward')
add('Footer / motion accessible names', 'Reduced motion follows your device setting', 'Resume ambient motion', 'Reduce motion')
add('Footer / step control accessible names', 'Return to entrance', 'Go to next room')
add('Floor plan / map accessible description', 'Interactive floor plan of seven portfolio rooms')
for project in data['projects']:
    if project.get('image'):
        add('Project screenshot / image alternative text / ' + project['title'], project['title'] + ' project screenshot')
add('JavaScript disabled / heading and description', 'Mahera Tasfee', 'Supply chain, operations, analytics, and digital systems.')
add('JavaScript disabled / links', 'Explore GitHub', 'Connect on LinkedIn')

section('Browser and social preview copy')
add('Browser title and social preview title', 'Mahera Tasfee | Supply Chain, Operations & Analytics')
add('Search description', 'Portfolio of Mahera Tasfee, a Supply Chain Management and Marketing BBA candidate focused on operations, planning, analytics, and practical digital systems.')
add('Search keywords', 'Mahera Tasfee', 'supply chain management', 'operations', 'demand planning', 'business analytics', 'Power BI', 'Dhaka')
add('Author and social site name', 'Mahera Tasfee')
add('Open Graph description', 'Turning operational complexity into decision-ready clarity through supply-chain thinking, analytics, and practical digital systems.')
add('Open Graph image alternative text', 'Mahera Tasfee — Supply chain, analytics, and digital systems')
add('Twitter preview description', 'Supply chain · analytics · digital systems')

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(0.65)
sec.left_margin = sec.right_margin = Inches(0.65)
for style_name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    s = doc.styles[style_name]
    s.font.name = 'Calibri'
    s.font.color.rgb = RGBColor(0, 0, 0)
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.06
doc.styles['Title'].font.size = Pt(25)
doc.styles['Title'].paragraph_format.space_after = Pt(12)
for style in doc.styles:
    for border in style.element.xpath('./w:pPr/w:pBdr'):
        border.getparent().remove(border)
doc.add_paragraph('Portfolio website copy inventory', style='Title')
doc.add_paragraph('Mahera Tasfee  |  5 October 2026')
doc.add_paragraph('Use this table to review the wording throughout the portfolio. Each position identifies where the current text appears. Reused exhibit copy is grouped with its shared locations; alternate control states and accessibility text are included. Project screenshot contents are separate visual assets.')

table = doc.add_table(rows=1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [Inches(2.35), Inches(4.85)]
for col, width in zip(table.columns, widths): col.width = width
table.rows[0].cells[0].text = 'Position'
table.rows[0].cells[1].text = 'Current text'
header = OxmlElement('w:tblHeader')
table.rows[0]._tr.get_or_add_trPr().append(header)

def shade(cell, color):
    el = OxmlElement('w:shd'); el.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(el)

def format_cell(cell, header=False, fill=None):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tcpr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement('w:tcMar')
    for edge, value in [('top', 95), ('bottom', 95), ('left', 125), ('right', 125)]:
        child = OxmlElement('w:' + edge); child.set(qn('w:w'), str(value)); child.set(qn('w:type'), 'dxa'); margins.append(child)
    tcpr.append(margins)
    if fill: shade(cell, fill)
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.06
        for run in p.runs:
            run.font.name = 'Calibri'; run.font.size = Pt(10.5)
            if header: run.bold = True; run.font.color.rgb = RGBColor(255, 255, 255)

for c in table.rows[0].cells: format_cell(c, header=True, fill='303640')
count = 0
for entry in rows:
    row = table.add_row()
    for cell, width in zip(row.cells, widths): cell.width = width
    pr = row._tr.get_or_add_trPr(); pr.append(OxmlElement('w:cantSplit'))
    if 'section' in entry:
        cell = row.cells[0].merge(row.cells[1]); cell.text = entry['section']
        format_cell(cell, fill='E6E9ED')
        cell.paragraphs[0].paragraph_format.keep_with_next = True
        for run in cell.paragraphs[0].runs: run.bold = True
    else:
        row.cells[0].text = entry['position']
        row.cells[1].text = entry['text']
        for cell in row.cells: format_cell(cell, fill='F7F8FA' if count % 2 else 'FFFFFF')
        count += 1

borders = OxmlElement('w:tblBorders')
for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
    el = OxmlElement('w:' + edge)
    for name, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]: el.set(qn('w:' + name), value)
    borders.append(el)
table._tbl.tblPr.append(borders)

# Page numbers are useful for a multi-page editing inventory.
p = sec.footer.paragraphs[0]; p.alignment = 2
run = p.add_run('Page '); run.font.size = Pt(9)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); p._p.append(field)
doc.core_properties.title = 'Portfolio website copy inventory'
doc.core_properties.subject = 'Current portfolio copy by position'
doc.core_properties.author = 'Mahera Tasfee'
output = ROOT / 'docs' / 'portfolio-copy-inventory.docx'
output.parent.mkdir(exist_ok=True)
doc.save(output)
(ROOT / 'docs' / 'portfolio-copy-inventory.json').write_text(json.dumps(rows, ensure_ascii=False, indent=2))
assert len(data['projects']) == 15
assert all(item['body'] in '\n'.join(r.get('text', '') for r in rows) for items in data['roomItems'].values() for item in items)
print(json.dumps({'path': str(output), 'copy_records': count, 'exhibits': 35, 'projects': 15}))
